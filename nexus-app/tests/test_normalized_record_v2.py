"""Tests for the B3 normalized_record v2 schema upgrade.

Three layers:

  1. `_build_normalized_record` unit tests — pin v2 field layout, profile +
     domain_profile placement, body_markdown placeholder null, and the
     backward-compatible behavior when `profile_dict` is None (JSON ingestion
     path).
  2. `_persist_normalized_ref` schema_version differentiation — record-typed
     refs persist NORMALIZED_RECORD_SCHEMA_VERSION, document-typed refs keep
     the legacy NORMALIZED_DOCUMENT_SCHEMA_VERSION.
  3. End-to-end: sample 1 xlsx → asserts both MinIO payload + PG row carry
     the v2 schema_version + domain_profile.
"""
from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from nexus_app import models, services
from nexus_app.config import Settings
from nexus_app.enums import (
    AuditEventType,
    JobStatus,
    NormalizedType,
)
from nexus_app.ingest.gateway import submit_file_bytes
from nexus_app.mineru import FakeMinerUAdapter
from nexus_app.pipeline.normalized_record_schema import (
    NORMALIZED_DOCUMENT_SCHEMA_VERSION,
    NORMALIZED_RECORD_SCHEMA_VERSION,
)
from nexus_app.pipeline.stages import _build_normalized_record
from nexus_app.schemas import DataSourceCreate
from nexus_app.storage import InMemoryObjectStorage
from nexus_app.worker.runner import execute_job

REPO_ROOT = Path(__file__).resolve().parents[2]
SAMPLE_JOB_DEMAND = REPO_ROOT / "docs/samples/1.（岗位需求）电子商务岗位招聘数据.xlsx"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# Schema constants
# ---------------------------------------------------------------------------


class TestSchemaConstants:
    def test_record_schema_version_pinned(self):
        # If this changes downstream consumers (B4 / B6 / B7) need to bump too.
        assert NORMALIZED_RECORD_SCHEMA_VERSION == "normalized-record.v2"

    def test_document_schema_version_unchanged(self):
        # Pipeline A is intentionally NOT bumped in B3 — keeping the
        # document chain stable while record evolves.
        assert NORMALIZED_DOCUMENT_SCHEMA_VERSION == "normalized-document-v1"

    def test_versions_are_distinct(self):
        assert NORMALIZED_RECORD_SCHEMA_VERSION != NORMALIZED_DOCUMENT_SCHEMA_VERSION


# ---------------------------------------------------------------------------
# _build_normalized_record — v2 payload layout
# ---------------------------------------------------------------------------


def _stub_raw_object():
    """Build an unsaved RawObject suitable for `_build_normalized_record` unit tests."""
    from nexus_app.enums import DataSourceType, RawObjectStatus
    obj = models.RawObject(
        data_source_id="ds-build",
        batch_id="batch-build",
        source_type=DataSourceType.FILE_UPLOAD,
        object_uri="s3://bucket/path/file.xlsx",
        checksum="cs-build",
        mime_type=XLSX_MIME,
        size_bytes=1,
        status=RawObjectStatus.RAW_PERSISTED,
        metadata_summary={"filename": "file.xlsx", "record_type": "generic"},
    )
    # _build_normalized_record reads raw_object.id into lineage; tests pass an
    # unsaved instance, so synthesize a stable id here.
    obj.id = "raw-build"
    return obj


def _profile_dict(record_type: str = "job_demand_dataset") -> dict:
    """Build a minimal ProfileDetectResult-shaped dict."""
    return {
        "record_type": record_type,
        "domain": "occupation",
        "domain_profile": "job_demand.v1",
        "analysis_model": None,
        "detector_version": "record-profile-detector.v1",
        "confidence": 0.95,
        "evidence": {
            "matched_headers": ["岗位名称", "城市"],
            "sheet_names": ["Sheet1"],
            "sample_row_count": 3,
            "matched_categories": [],
            "matched_code_prefixes": [],
        },
    }


class TestBuildNormalizedRecordWithProfile:
    def test_schema_version_is_v2(self):
        payload = _build_normalized_record(
            _stub_raw_object(), {"records": []}, profile_dict=_profile_dict()
        )
        assert payload["schema_version"] == NORMALIZED_RECORD_SCHEMA_VERSION

    def test_top_level_domain_profile_set_from_profile(self):
        payload = _build_normalized_record(
            _stub_raw_object(), {"records": []}, profile_dict=_profile_dict()
        )
        assert payload["domain_profile"] == "job_demand.v1"

    def test_record_type_uses_profile_record_type(self):
        # Profile detection wins over raw_object.metadata_summary fallback.
        payload = _build_normalized_record(
            _stub_raw_object(), {}, profile_dict=_profile_dict("occupational_ability_analysis")
        )
        assert payload["record_type"] == "occupational_ability_analysis"

    def test_body_markdown_placeholders_present_and_null(self):
        # B5 will populate these; v2 carries them as nulls so downstream
        # readers don't need to special-case "missing field" vs "null".
        payload = _build_normalized_record(
            _stub_raw_object(), {}, profile_dict=_profile_dict()
        )
        assert "body_markdown" in payload
        assert payload["body_markdown"] is None
        assert "body_markdown_meta" in payload
        assert payload["body_markdown_meta"] is None

    def test_profile_written_to_top_level_and_metadata_mirror(self):
        profile = _profile_dict()
        payload = _build_normalized_record(
            _stub_raw_object(), {}, profile_dict=profile
        )
        assert payload["profile"] == profile
        # metadata mirrors profile + domain_profile so NormalizedAssetRef.
        # metadata_summary can be filtered on the PG row without a MinIO hit.
        assert payload["metadata"]["profile"] == profile
        assert payload["metadata"]["domain_profile"] == "job_demand.v1"

    def test_record_body_projected_to_contract_shape_when_profile_present(self):
        # B3.5 (record_body_adapter) projects the ParsedWorkbook dump into
        # `{dataset, records}` for job_demand.v1, so the writer doesn't have
        # to re-derive the shape. A dict that doesn't look like a workbook
        # still passes through the projector — it just yields an empty
        # records list (no sheets means no records).
        raw = {"sheets": []}
        payload = _build_normalized_record(
            _stub_raw_object(), raw, profile_dict=_profile_dict()
        )
        # job_demand.v1 path always emits the contract envelope.
        assert "dataset" in payload["record_body"]
        assert "records" in payload["record_body"]
        assert payload["record_body"]["records"] == []

    def test_governance_quality_lineage_present(self):
        # Acceptance criterion (implementation plan §三 B3): all of these
        # must be non-null after _build_normalized_record runs.
        payload = _build_normalized_record(
            _stub_raw_object(), {}, profile_dict=_profile_dict()
        )
        assert payload["governance"] is not None
        assert payload["quality"] is not None
        assert payload["lineage"] is not None
        assert payload["lineage"]["raw_object_id"] is not None


class TestBuildNormalizedRecordWithoutProfile:
    """Backward-compatible path: JSON ingestion (crawler / database / webhook)
    that doesn't run profile_detect must still produce a valid v2 payload.
    """

    def test_v2_schema_version_still_used(self):
        # The schema bump applies to ALL record_type normalize_record output —
        # not just the profile-detected paths.
        payload = _build_normalized_record(
            _stub_raw_object(), {"k": "v"}, profile_dict=None
        )
        assert payload["schema_version"] == NORMALIZED_RECORD_SCHEMA_VERSION

    def test_domain_profile_is_none_when_no_profile(self):
        payload = _build_normalized_record(
            _stub_raw_object(), {"k": "v"}, profile_dict=None
        )
        assert payload["domain_profile"] is None

    def test_record_type_falls_back_to_raw_object_metadata(self):
        raw_object = _stub_raw_object()
        raw_object.metadata_summary["record_type"] = "legacy_crawler_record"
        payload = _build_normalized_record(raw_object, {}, profile_dict=None)
        assert payload["record_type"] == "legacy_crawler_record"

    def test_profile_key_absent_when_no_profile(self):
        # Don't synthesize an empty profile — leave it absent so B7 governance
        # can disjoin on "structured_parse path" vs "JSON path".
        payload = _build_normalized_record(
            _stub_raw_object(), {}, profile_dict=None
        )
        assert "profile" not in payload
        assert "profile" not in payload["metadata"]
        assert "domain_profile" not in payload["metadata"]

    def test_body_markdown_placeholders_still_present(self):
        payload = _build_normalized_record(
            _stub_raw_object(), {}, profile_dict=None
        )
        assert payload["body_markdown"] is None
        assert payload["body_markdown_meta"] is None


# ---------------------------------------------------------------------------
# _persist_normalized_ref — schema_version on the PG row
# ---------------------------------------------------------------------------


@pytest.mark.skipif(not SAMPLE_JOB_DEMAND.exists(), reason="sample missing")
class TestPersistedSchemaVersion:
    """Worker-driven integration: persist a record-typed NormalizedAssetRef
    end-to-end and verify both the MinIO JSON and the PG row carry the v2
    schema_version. Document-typed refs untouched in B3 are covered by the
    existing `test_week2_pipeline` suite (no regression already proven by
    the full-suite run).
    """

    @pytest.fixture
    def ref(self, session):
        storage = InMemoryObjectStorage()
        source = services.create_data_source(
            session,
            DataSourceCreate(code="b3-sample1", name="b3-sample1", source_type="file_upload"),
        )
        accepted = submit_file_bytes(
            session,
            data_source_id=source.id,
            idempotency_key="b3-sample1-key",
            content=SAMPLE_JOB_DEMAND.read_bytes(),
            filename=SAMPLE_JOB_DEMAND.name,
            content_type=XLSX_MIME,
            storage=storage,
            settings=Settings(pipeline_b_xlsx_enabled=True),
            trace_id="trace-b3",
        )
        session.refresh(accepted.job)
        execute_job(
            accepted.job,
            session,
            storage,
            FakeMinerUAdapter(),
            Settings(pipeline_b_xlsx_enabled=True),
        )
        session.refresh(accepted.job)
        assert accepted.job.status == JobStatus.SUCCEEDED, (
            f"job failed: {accepted.job.failure_reason}"
        )

        ref = next(iter(session.scalars(select(models.NormalizedAssetRef))))
        return ref, storage

    def test_pg_row_carries_v2_schema_version(self, ref, session):
        r, _ = ref
        assert r.schema_version == NORMALIZED_RECORD_SCHEMA_VERSION
        assert r.normalized_type == NormalizedType.RECORD

    def test_minio_payload_carries_v2_schema_version(self, ref, session):
        r, storage = ref
        key = r.object_uri.split("/", 3)[-1]
        payload = json.loads(storage.get_bytes(key).decode("utf-8"))
        assert payload["schema_version"] == NORMALIZED_RECORD_SCHEMA_VERSION

    def test_minio_payload_carries_top_level_domain_profile(self, ref, session):
        r, storage = ref
        key = r.object_uri.split("/", 3)[-1]
        payload = json.loads(storage.get_bytes(key).decode("utf-8"))
        # B2 wrote profile.domain_profile; B3 mirrors it to the top level
        # so consumers can route without parsing the profile dict.
        assert payload["domain_profile"] == "job_demand.v1"
        # body_markdown placeholders ride along even before B5 fills them.
        assert payload["body_markdown"] is None
        assert payload["body_markdown_meta"] is None

    def test_metadata_summary_includes_domain_profile_mirror(self, ref):
        r, _ = ref
        assert r.metadata_summary.get("domain_profile") == "job_demand.v1"
        # And the full profile dict is still there from B2.3.
        assert r.metadata_summary.get("profile", {}).get("record_type") == "job_demand_dataset"

    def test_v2_required_fields_non_empty(self, ref, session):
        # Implementation plan §三 B3 acceptance: "profile / quality / lineage
        # fields non-empty after B1+B2+B3 end-to-end".
        r, storage = ref
        key = r.object_uri.split("/", 3)[-1]
        payload = json.loads(storage.get_bytes(key).decode("utf-8"))
        assert payload["profile"]
        assert payload["quality"] is not None
        assert payload["lineage"]
        assert payload["lineage"]["raw_object_id"]

    def test_governance_result_target_still_points_at_normalized_ref(self, ref, session):
        # Architectural invariant from CLAUDE.md (and acceptance criterion):
        # `governance_result.target = normalized_asset_ref`. B3 doesn't add
        # any back-pointer that would invert this — we verify by checking
        # the GovernanceResult model still references normalized_ref_id
        # rather than the NormalizedAssetRef referencing a governance_result.
        # No GovernanceResult is created in the test (no active rules), but
        # the schema must still hold the one-direction invariant.
        assert hasattr(models.GovernanceResult, "normalized_ref_id")
        assert not hasattr(models.NormalizedAssetRef, "governance_result_id")
