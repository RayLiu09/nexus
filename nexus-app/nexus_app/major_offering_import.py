"""Bounded Excel import into the direct offering domain, without pipelines."""

from hashlib import sha256
from io import BytesIO
from pathlib import PurePath
import re
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from nexus_app import models
from nexus_app.audit import write_audit
from nexus_app.data_sync.adapters.moe_major_offering import PROVINCES
from nexus_app.data_sync.major_offering_intake import insert_offering
from nexus_app.data_sync.runtime import SyncRuntimeError
from nexus_app.domain_normalize.administrative_division import normalize_province_name
from nexus_app.enums import AuditEventType

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
HEADERS = ("年份", "省份", "专业代码", "专业名称", "学校标识码", "学校名称", "年限", "备注")
FIELDS = ("year", "province_name", "major_code", "major_name", "school_code", "school_name", "study_duration", "remark")
ALIASES = {"省份/地区": "省份", "学校标识": "学校标识码", "修业年限": "年限"}


class ImportInvalid(ValueError):
    pass


class ImportConflict(ValueError):
    pass


class OfferingInput(BaseModel):
    model_config = ConfigDict(str_strip_whitespace=True)
    year: int = Field(ge=2013, le=2035)
    province_name: str = Field(min_length=1, max_length=128)
    major_code: str = Field(pattern=r"^[0-9]{6}$")
    major_name: str = Field(min_length=1, max_length=200)
    school_code: str | None = Field(default=None, max_length=64)
    school_name: str = Field(min_length=1, max_length=256)
    study_duration: str = Field(min_length=1, max_length=64)
    remark: str | None = Field(default=None, max_length=2000)


def _text(cell):
    value = cell.value
    if cell.data_type == "f":
        raise ImportInvalid("不支持公式，请将公式转换为值后导入")
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise ImportInvalid("请使用文本或数字")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and re.fullmatch(r"0{2,20}", cell.number_format or ""):
        return str(value).zfill(len(cell.number_format))
    return str(value).strip() or None


def parse_excel(content: bytes) -> list[tuple[int, dict]]:
    if len(content) > MAX_BYTES:
        raise ImportInvalid("文件不能超过 10 MiB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ImportInvalid("工作簿展开后过大，请拆分文件")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except ImportInvalid:
        raise
    except (BadZipFile, InvalidFileException, ValueError, KeyError, OSError, SyntaxError) as exc:
        raise ImportInvalid("无法读取 Excel，请使用有效的 .xlsx 文件") from exc
    try:
        sheet = workbook.active
        if sheet is None or (sheet.max_column or 0) > 32:
            raise ImportInvalid("请使用模板，工作表最多支持 32 列")
        rows = sheet.iter_rows()
        header = next(rows, ())
        titles = [ALIASES.get(str(c.value).strip(), str(c.value).strip()) if c.value is not None else "" for c in header]
        present = [title for title in titles if title in HEADERS]
        required = set(HEADERS) - {"学校标识码", "备注"}
        missing = required - set(present)
        if missing:
            raise ImportInvalid("缺少表头：" + "、".join(sorted(missing)))
        if len(present) != len(set(present)):
            raise ImportInvalid("表头重复，请使用模板")
        columns = {FIELDS[HEADERS.index(title)]: i for i, title in enumerate(titles) if title in HEADERS}
        result = []
        for number, cells in enumerate(rows, start=2):
            if number > MAX_ROWS + 1:
                raise ImportInvalid("工作表最多支持 10,000 行数据，请拆分文件")
            if all(cell.value is None for cell in cells):
                continue
            try:
                values = {key: _text(cells[index]) for key, index in columns.items()}
                data = OfferingInput.model_validate(values).model_dump()
                province = normalize_province_name(data["province_name"])
                code = next((code for code, name in PROVINCES.items() if name == province), None)
                if code is None:
                    raise ImportInvalid("省份不在支持的省份/地区范围内")
                data.update(province_name=province, province_code=code)
                result.append((number, data))
            except ValidationError as exc:
                names = [HEADERS[FIELDS.index(error["loc"][0])] for error in exc.errors()]
                raise ImportInvalid(f"第 {number} 行：{'、'.join(dict.fromkeys(names))}缺失或格式不正确") from exc
            except ImportInvalid as exc:
                raise ImportInvalid(f"第 {number} 行：{exc}") from exc
        if not result:
            raise ImportInvalid("工作表没有可导入的数据")
        return result
    finally:
        workbook.close()


def import_excel(session, *, content, filename, actor_id, idempotency_key, trace_id):
    digest = sha256(content).hexdigest()

    def replay():
        previous = session.scalar(select(models.MajorOfferingImport).where(
            models.MajorOfferingImport.created_by == actor_id,
            models.MajorOfferingImport.idempotency_key == idempotency_key,
        ))
        if previous is not None and previous.request_hash != digest:
            raise ImportConflict("该幂等键已用于其他文件，请重新选择文件后提交")
        return previous

    existing = replay()
    if existing is not None:
        return {"import_id": existing.id, **existing.result_summary}
    rows = parse_excel(content)
    receipt = models.MajorOfferingImport(
        created_by=actor_id, idempotency_key=idempotency_key,
        request_hash=digest, filename=PurePath(filename).name[:256], result_summary={},
    )
    session.add(receipt)
    try:
        session.flush()
    except IntegrityError:
        session.rollback()
        existing = replay()
        if existing is None:
            raise
        return {"import_id": existing.id, **existing.result_summary}
    accepted = 0
    try:
        for number, values in rows:
            try:
                accepted += insert_offering(session, {
                    **values, "import_id": receipt.id, "provider_code": "excel_import", "source_url": "",
                })
            except SyncRuntimeError as exc:
                raise ImportInvalid(f"第 {number} 行：学校名称与标识码对应关系冲突，请核对") from exc
        summary = {"total_count": len(rows), "imported_count": accepted, "skipped_count": len(rows) - accepted}
        receipt.result_summary = summary
        write_audit(
            session, AuditEventType.MAJOR_OFFERINGS_IMPORTED, "major_offering_import", receipt.id,
            trace_id, {"filename": receipt.filename, **summary}, actor_type="user", actor_id=actor_id,
        )
        session.commit()
        return {"import_id": receipt.id, **summary}
    except Exception:
        session.rollback()
        raise
