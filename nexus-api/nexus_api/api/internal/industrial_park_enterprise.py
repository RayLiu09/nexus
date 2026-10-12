"""Asset Center catalog APIs for industrial parks and enterprises."""

from __future__ import annotations

from tempfile import SpooledTemporaryFile

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import Text, cast, func, or_, select
from sqlalchemy.orm import Session

from nexus_api import schemas
from nexus_api.dependencies import Pagination, pagination_params
from nexus_api.responses import list_response
from nexus_app import models
from nexus_app.database import get_db

router = APIRouter()


def _excel_text(value: str | None) -> str:
    text = (value or "").translate({code: None for code in range(32) if code not in (9, 10, 13)})
    return ("'" + text if text.lstrip().startswith("=") else text)[:32767]


def _park_payload(row: models.IndustrialPark, enterprise_count: int) -> dict:
    return {
        "id": row.id,
        "park_name": row.park_name,
        "industry_tags": row.industry_tags or [],
        "introduction": row.introduction,
        "enterprise_count": enterprise_count,
        "source_url": row.source_url,
        "created_at": row.created_at.isoformat() if row.created_at else None,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
    }


def _enterprise_payload(row: models.Enterprise, park: models.IndustrialPark | None) -> dict:
    return {
        "id": row.id,
        "enterprise_name": row.enterprise_name,
        "business_scope": row.business_scope,
        "industrial_park_id": row.industrial_park_id,
        "park_name": park.park_name if park else None,
        "source_url": row.source_url,
        "created_at": row.created_at.isoformat() if row.created_at else None,
        "updated_at": row.updated_at.isoformat() if row.updated_at else None,
    }


@router.get("/industrial-parks", response_model=schemas.ListResponse[dict])
def list_industrial_parks(
    request: Request,
    q: str | None = Query(None, min_length=1, max_length=256),
    industry: str | None = Query(None, min_length=1, max_length=128),
    pagination: Pagination = Depends(pagination_params),
    session: Session = Depends(get_db),
):
    clauses = []
    if q:
        clauses.append(or_(models.IndustrialPark.park_name.ilike(f"%{q}%"), models.IndustrialPark.introduction.ilike(f"%{q}%")))
    if industry:
        clauses.append(cast(models.IndustrialPark.industry_tags, Text).ilike(f"%{industry}%"))
    count = session.scalar(select(func.count(models.IndustrialPark.id)).where(*clauses)) or 0
    rows = session.scalars(
        select(models.IndustrialPark)
        .where(*clauses)
        .order_by(models.IndustrialPark.park_name, models.IndustrialPark.id)
        .offset(pagination.offset).limit(pagination.limit)
    ).all()
    counts = dict(session.execute(
        select(models.Enterprise.industrial_park_id, func.count(models.Enterprise.id))
        .where(models.Enterprise.industrial_park_id.in_([row.id for row in rows]))
        .group_by(models.Enterprise.industrial_park_id)
    ).all()) if rows else {}
    return list_response([_park_payload(row, counts.get(row.id, 0)) for row in rows], request, page=pagination.page, page_size=pagination.page_size, total=count)


@router.get("/enterprises", response_model=schemas.ListResponse[dict])
def list_enterprises(
    request: Request,
    q: str | None = Query(None, min_length=1, max_length=256),
    park_id: str | None = Query(None, max_length=36),
    pagination: Pagination = Depends(pagination_params),
    session: Session = Depends(get_db),
):
    clauses = []
    if q:
        pattern = f"%{q}%"
        clauses.append(or_(models.Enterprise.enterprise_name.ilike(pattern), models.Enterprise.business_scope.ilike(pattern), models.IndustrialPark.park_name.ilike(pattern)))
    if park_id:
        clauses.append(models.Enterprise.industrial_park_id == park_id)
    base = select(models.Enterprise).outerjoin(models.IndustrialPark).where(*clauses)
    count = session.scalar(select(func.count(models.Enterprise.id)).select_from(models.Enterprise).outerjoin(models.IndustrialPark).where(*clauses)) or 0
    rows = session.scalars(base.order_by(models.Enterprise.enterprise_name, models.Enterprise.id).offset(pagination.offset).limit(pagination.limit)).all()
    return list_response([_enterprise_payload(row, row.industrial_park) for row in rows], request, page=pagination.page, page_size=pagination.page_size, total=count)


@router.get("/industrial-parks/export")
def export_industrial_parks(session: Session = Depends(get_db)):
    workbook = Workbook(write_only=True); sheet = workbook.create_sheet("产业园区")
    sheet.append(["园区名称", "行业标签", "园区介绍", "企业数量", "来源链接"])
    for row in session.scalars(select(models.IndustrialPark).order_by(models.IndustrialPark.park_name)).yield_per(1000):
        count = session.scalar(select(func.count(models.Enterprise.id)).where(models.Enterprise.industrial_park_id == row.id)) or 0
        sheet.append([_excel_text(row.park_name), _excel_text("、".join(row.industry_tags or [])), _excel_text(row.introduction), count, _excel_text(row.source_url)])
    return _xlsx_response(workbook, "industrial-parks.xlsx", "产业园区数据.xlsx")


@router.get("/enterprises/export")
def export_enterprises(session: Session = Depends(get_db)):
    workbook = Workbook(write_only=True); sheet = workbook.create_sheet("企业")
    sheet.append(["企业名称", "经营范围", "所属产业园区", "来源链接"])
    rows = session.execute(select(models.Enterprise, models.IndustrialPark.park_name).outerjoin(models.IndustrialPark).order_by(models.Enterprise.enterprise_name)).yield_per(1000)
    for row, park_name in rows:
        sheet.append([_excel_text(row.enterprise_name), _excel_text(row.business_scope), _excel_text(park_name), _excel_text(row.source_url)])
    return _xlsx_response(workbook, "enterprises.xlsx", "企业数据.xlsx")


def _xlsx_response(workbook: Workbook, ascii_name: str, display_name: str) -> StreamingResponse:
    output = SpooledTemporaryFile(max_size=8 * 1024 * 1024)
    workbook.save(output); output.seek(0)
    def chunks():
        try:
            while chunk := output.read(65536): yield chunk
        finally: output.close()
    from urllib.parse import quote
    return StreamingResponse(chunks(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={ascii_name}; filename*=UTF-8''{quote(display_name)}", "Cache-Control": "no-store"})
