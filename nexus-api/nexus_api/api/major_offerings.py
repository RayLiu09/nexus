"""Offering reads plus role-checked Console Excel import."""

from datetime import datetime
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from nexus_api import schemas
from nexus_api.dependencies import Pagination, pagination_params, require_api_caller, require_idempotency_key, require_user
from nexus_api.responses import list_response, response
from nexus_app import major_offerings, models
from nexus_app.audit import write_audit
from nexus_app.database import get_db
from nexus_app.enums import AuditEventType, UserRole


class MajorOfferingRead(BaseModel):
    id: str
    year: int
    province_code: str
    province_name: str
    major_code: str
    major_name: str
    school_code: str | None
    school_name: str
    study_duration: str
    remark: str | None
    source_url: str
    provider_code: str
    sync_run_id: str | None
    import_id: str | None
    fetched_at: datetime
    created_at: datetime


class MajorOfferingGroup(BaseModel):
    year: int | None = None
    province_code: str | None = None
    province_name: str | None = None
    major_code: str | None = None
    major_name: str | None = None
    record_count: int
    institution_count: int


def retired_endpoint():
    raise HTTPException(410, "专业布点旧管道接口已退役，请使用 /v1/major-offerings")


def filters(
    year: int | None = Query(None, ge=2013, le=2200, description="登记年份，精确匹配；查询已入库数据，不触发同步"),
    province_name: str | None = Query(None, max_length=128, description="省份名称，行政区名称归一化后精确匹配"),
    province_code: str | None = Query(None, max_length=16, description="教育部省份编码，区分大小写，精确匹配，如 BJ、HeB"),
    major_name: str | None = Query(None, max_length=200, description="专业名称，字面子串匹配"),
    major_code: str | None = Query(None, max_length=64, description="专业代码，精确匹配"),
    school_name: str | None = Query(None, max_length=256, description="学校名称，字面子串匹配"),
    school_code: str | None = Query(None, max_length=64, description="学校标识，精确匹配；来源可能未提供该字段"),
):
    return {
        "year": year,
        "province_name": province_name,
        "province_code": province_code,
        "major_name": major_name,
        "major_code": major_code,
        "school_name": school_name,
        "school_code": school_code,
    }


def _read_audit(session, request, actor, total, action="list"):
    write_audit(
        session,
        AuditEventType.OPEN_RECORD_ASSETS_ACCESSED,
        "major_offering_record",
        action,
        str(request.state.trace_id),
        {"result_count": total},
        actor_type="api_caller" if isinstance(actor, models.ApiCaller) else "user",
        actor_id=actor.id,
    )
    session.commit()


def make_router(prefix, auth):
    router = APIRouter(prefix=prefix, tags=["major-offerings"], dependencies=[Depends(auth)])

    # Binary downloads serve the Console and canonical business API; JSON Agent
    # tools continue to use the /open/v1 list and aggregation endpoints.
    if not prefix.startswith("/open/v1"):
        @router.get("/export")
        def export_major_offerings(
            request: Request,
            criteria: dict = Depends(filters),
            actor=Depends(auth),
            session: Session = Depends(get_db),
        ):
            """导出当前筛选条件下全部专业布点为 Excel，不受列表分页限制。"""
            from openpyxl import Workbook
            from nexus_api.api.internal.industrial_park_enterprise import _excel_text, _xlsx_response

            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("专业布点")
            sheet.append(["年份", "省份", "专业代码", "专业名称", "学校标识码", "学校名称", "年限", "备注"])
            count = 0
            stmt = major_offerings.ordered(major_offerings.statement(**criteria))
            for row in session.scalars(stmt).yield_per(1000):
                sheet.append([row.year, *(_excel_text(getattr(row, key)) for key in (
                    "province_name", "major_code", "major_name", "school_code",
                    "school_name", "study_duration", "remark",
                ))])
                count += 1
            result = _xlsx_response(workbook, "major-offerings.xlsx", "专业布点.xlsx")
            _read_audit(session, request, actor, count, action="export")
            return result

    @router.get("", response_model=schemas.ListResponse[MajorOfferingRead],
                operation_id="list_major_offerings" if prefix.startswith("/open/v1") else None)
    def list_major_offerings(
        request: Request,
        criteria: dict = Depends(filters),
        pagination: Pagination = Depends(pagination_params),
        actor=Depends(auth),
        session: Session = Depends(get_db),
    ):
        """查询教育部政务服务平台同步及控制台 Excel 导入的专业布点数据，支持年份、省份、专业和学校组合过滤与分页。

        返回省份、专业代码/名称、学校标识（可空）、学校名称、修业年限、备注及来源信息。
        同一年份、同一学校、同一专业去重；同步直接入库，无 AI 治理或资产版本准入。
        sync_run_id 和 import_id 分别标识 API 同步和 Excel 导入来源，二者仅有一个非空。
        仅表示已入库覆盖，不代表全国总体，不读取已废弃的专业布点 Pipeline B 投影。
        """
        rows, total = major_offerings.list_records(
            session, offset=pagination.offset, limit=pagination.limit, **criteria
        )
        _read_audit(session, request, actor, total)
        return list_response(
            rows, request, page=pagination.page, page_size=pagination.page_size, total=total
        )

    @router.get("/aggregate", response_model=schemas.ListResponse[MajorOfferingGroup],
                response_model_exclude_unset=True,
                operation_id="aggregate_major_offerings" if prefix.startswith("/open/v1") else None)
    def aggregate_major_offerings(
        request: Request,
        criteria: dict = Depends(filters),
        group_by: list[Literal["year", "province_code", "province_name", "major_code", "major_name"]] = Query(
            ["year", "province_name", "major_code", "major_name"], min_length=1, max_length=5,
            description="分组维度（不能重复），使用重复 query key 传数组；不传时按年份、省份名称、专业代码/名称分组"),
        pagination: Pagination = Depends(pagination_params),
        actor=Depends(auth),
        session: Session = Depends(get_db),
    ):
        """对教育部 API 同步及 Excel 导入专业布点按白名单维度统计全量筛选结果，再对分组结果分页。

        返回选定维度、record_count（登记记录数）、institution_count（学校标识优先、缺失时按归一化学校名称去重的学校数）。
        meta.total 为分组数；跨年分组时登记记录数可大于学校数。
        数据仅代表已入库覆盖，不使用专业简介或培养方案推导布点，不触发同步或 AI 治理。
        """
        model = models.MajorOfferingRecord
        allowed = {
            name: getattr(model, name)
            for name in ("year", "province_code", "province_name", "major_code", "major_name")
        }
        if not group_by or len(group_by) != len(set(group_by)) or set(group_by) - set(allowed):
            raise HTTPException(422, "invalid group_by")
        cols = [allowed[name] for name in group_by]
        # Name is the fallback school identity when no Ministry identifier exists.
        identity = func.coalesce(model.school_code, model.school_name_key)
        stmt = (
            major_offerings.statement(**criteria)
            .with_only_columns(
                *cols,
                func.count(model.id).label("record_count"),
                func.count(func.distinct(identity)).label("institution_count"),
            )
            .group_by(*cols)
        )
        total = session.scalar(select(func.count()).select_from(stmt.subquery())) or 0
        rows = (
            session.execute(stmt.order_by(*cols).offset(pagination.offset).limit(pagination.limit))
            .mappings()
            .all()
        )
        _read_audit(session, request, actor, total)
        return list_response(
            [dict(row) for row in rows],
            request,
            page=pagination.page,
            page_size=pagination.page_size,
            total=total,
        )

    return router


internal_router = make_router("/major-offerings", require_user)
open_router = make_router("/open/v1/major-offerings", require_api_caller)
business_router = make_router("/v1/major-offerings", require_api_caller)


class MajorOfferingImportRead(BaseModel):
    import_id: str
    total_count: int
    imported_count: int
    skipped_count: int


def require_import_user(user: models.UserAccount = Depends(require_user)):
    if user.role not in (UserRole.PLATFORM_DATA_ADMIN, UserRole.BUSINESS_EXPERT):
        raise HTTPException(403, "仅业务专家或平台数据管理员可导入布点数据")
    return user


@internal_router.get("/import-template")
def major_offering_import_template():
    from openpyxl import Workbook
    from nexus_app.major_offering_import import HEADERS
    from nexus_api.api.internal.industrial_park_enterprise import _xlsx_response

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "专业布点"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    for column, width in zip("ABCDEFGH", (12, 20, 16, 24, 20, 36, 12, 40)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, max_row=1001, max_col=8):
        for cell in row:
            cell.number_format = "@"
    return _xlsx_response(workbook, "major-offerings-template.xlsx", "专业布点导入模板.xlsx")


@internal_router.post("/import", response_model=schemas.ApiResponse[MajorOfferingImportRead])
def import_major_offerings(
    request: Request,
    file: UploadFile = File(...),
    user=Depends(require_import_user),
    idempotency_key: str = Depends(require_idempotency_key),
    session: Session = Depends(get_db),
):
    """控制台用户导入 Excel，整文件事务校验，重复事实跳过，不触发 AI 治理。"""
    from nexus_app.major_offering_import import MAX_BYTES, ImportConflict, ImportInvalid, import_excel

    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(422, "仅支持 .xlsx 文件，请下载模板填写")
    content = file.file.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise HTTPException(413, "文件不能超过 10 MiB")
    try:
        result = import_excel(
            session, content=content, filename=file.filename, actor_id=user.id,
            idempotency_key=idempotency_key, trace_id=str(request.state.trace_id),
        )
    except ImportConflict as exc:
        raise HTTPException(409, str(exc)) from exc
    except ImportInvalid as exc:
        raise HTTPException(422, str(exc)) from exc
    return response(result, request)
