from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from nexus_app import models


def test_raw_jobs_list_reads_raw_intake_fields_only(app, session):
    session.add_all(
        [
            models.RawJob(
                id="raw-job-1",
                title_raw="数据分析师",
                company_name_raw="示例科技",
                address_raw="上海市浦东新区",
                salary_raw="15-25K",
                experience_raw="3-5年",
                degree_raw="本科",
                responsibilities_raw="负责经营数据分析",
                company_industry_raw="互联网",
                company_scale_raw="100-499人",
            ),
            models.RawJob(id="raw-job-2", title_raw="前端工程师", company_name_raw="另一家公司"),
        ]
    )
    session.commit()

    with TestClient(app) as client:
        response = client.get("/internal/v1/raw-jobs", params={"q": "数据分析师"})

    assert response.status_code == 200
    body = response.json()
    assert body["meta"]["total"] == 1
    assert body["data"] == [
        {
            "id": "raw-job-1",
            "job_title": "数据分析师",
            "job_responsibilities": "负责经营数据分析",
            "experience_requirement": "3-5年",
            "education": "本科",
            "salary_range": "15-25K",
            "industry": "互联网",
            "company_name": "示例科技",
            "company_size": "100-499人",
            "address": "上海市浦东新区",
            "source_url": None,
            "collected_at": None,
            "created_at": body["data"][0]["created_at"],
        }
    ]


def test_raw_jobs_export_includes_every_matching_row_and_business_columns(app, session):
    session.add_all([
        models.RawJob(
            id=f"raw-job-{index}",
            title_raw=f"数据岗位{index}",
            company_industry_raw="互联网",
            responsibilities_raw="分析业务数据",
            company_name_raw="示例科技",
        )
        for index in range(25)
    ])
    session.add(models.RawJob(
        id="raw-job-other", title_raw="其它岗位", company_industry_raw="制造业"
    ))
    session.commit()

    with TestClient(app) as client:
        response = client.get("/internal/v1/raw-jobs/export", params={"industry": "互联网"})

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    sheet = load_workbook(BytesIO(response.content), read_only=True).active
    rows = list(sheet.values)
    assert rows[0] == (
        "岗位名称", "岗位职责要求", "经验要求", "学历", "薪资范围",
        "所在行业", "公司名称", "规模", "地址",
    )
    assert len(rows) == 26
    assert {row[0] for row in rows[1:]} == {f"数据岗位{index}" for index in range(25)}
    assert all(row[5] == "互联网" for row in rows[1:])


def test_raw_jobs_export_keeps_untrusted_text_as_excel_string(app, session):
    session.add(models.RawJob(
        id="raw-job-formula", title_raw="=1+1", description_raw="=HYPERLINK(\"x\")"
    ))
    session.commit()

    with TestClient(app) as client:
        response = client.get("/internal/v1/raw-jobs/export")

    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content), read_only=True).active
    cells = list(sheet.rows)[1]
    assert cells[0].value == "'=1+1"
    assert cells[0].data_type == "s"
    assert cells[1].value == "'=HYPERLINK(\"x\")"
