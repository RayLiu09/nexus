from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import func, select

from nexus_app import models
from nexus_app.enums import AuditEventType, UserRole
from nexus_app.major_offering_import import HEADERS, parse_excel
from test_major_offerings import seed

PATH = "/internal/v1/major-offerings/import"
ROW = [2026, "云南", "530701", "电子商务", "00123", "云南测试学院", "3", None]


def excel(rows, headers=HEADERS):
    workbook = Workbook()
    workbook.active.append(headers)
    for row in rows:
        workbook.active.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def upload(client, content, key="import-1", filename="布点.xlsx"):
    return client.post(PATH, files={"file": (filename, content)}, headers={"Idempotency-Key": key})


def count(session, model):
    return session.scalar(select(func.count()).select_from(model))


def test_import_deduplicates_api_and_excel_and_updates_asset_center_counts(app, session):
    seed(session)
    client = TestClient(app)
    counts_before = {model: count(session, model) for model in (
        models.DataSyncRun, models.Asset, models.NormalizedAssetRef, models.GovernanceResult, models.Job,
    )}
    content = excel([
        [2026, "北京", "530701", "电子商务", None, "学校甲", "5", "不覆盖已有数据"],
        ROW, ROW, [2025, *ROW[1:]], [2026, "云南省", "530702", "跨境电子商务", *ROW[4:]],
    ])
    result = upload(client, content)
    assert result.status_code == 200, result.text
    summary = result.json()["data"]
    assert {k: summary[k] for k in ("total_count", "imported_count", "skipped_count")} == {
        "total_count": 5, "imported_count": 4, "skipped_count": 1,
    }
    assert count(session, models.MajorOfferingRecord) == 7
    assert count(session, models.MajorOfferingImport) == 1
    assert all(count(session, model) == before for model, before in counts_before.items())
    assert session.get(models.MajorOfferingRecord, "1").study_duration == "3"
    listed = client.get("/v1/major-offerings", params={"province_name": "云南"}).json()
    assert listed["meta"]["total"] == 3
    assert all(row["sync_run_id"] is None and row["import_id"] == summary["import_id"] for row in listed["data"])
    assert all(row["school_code"] == "00123" and row["province_code"] == "YN" for row in listed["data"])
    stats = client.get("/internal/v1/asset-center/counts").json()["data"]["counts"]
    assert stats["major/distributions"] == client.get("/v1/major-offerings").json()["meta"]["total"] == 7
    audits = list(session.scalars(select(models.AuditLog).where(models.AuditLog.event_type == AuditEventType.MAJOR_OFFERINGS_IMPORTED)))
    assert len(audits) == 1 and audits[0].actor_id and audits[0].trace_id
    assert audits[0].summary["imported_count"] == 4
    assert upload(client, content).json()["data"] == summary
    assert upload(client, excel([ROW])).status_code == 409
    assert upload(client, content, key="import-2").json()["data"]["skipped_count"] == 5
    assert count(session, models.MajorOfferingRecord) == 7
    facts = client.get("/v1/major-offerings", params={"year": 2026, "school_name": "学校甲"}).json()["data"]
    assert {r["study_duration"] for r in facts} == {"3", "5"}
    assert client.get("/v1/major-offerings/aggregate", params={"year": 2026, "school_name": "学校甲", "group_by": "year"}).json()["data"] == [
        {"year": 2026, "record_count": 2, "institution_count": 1},
    ]


def test_code_identity_and_transactional_conflict(app, session):
    client = TestClient(app)
    assert upload(client, excel([ROW])).status_code == 200
    renamed = [*ROW[:5], "别名学院", *ROW[6:]]
    assert upload(client, excel([renamed]), key="alias").json()["data"]["skipped_count"] == 1
    conflict = [*ROW[:4], "00999", *ROW[5:]]
    new = [*ROW[:4], "00888", "新学校", *ROW[6:]]
    response = upload(client, excel([new, conflict]), key="conflict")
    assert response.status_code == 422 and "第 3 行" in response.text
    assert count(session, models.MajorOfferingRecord) == 1
    assert count(session, models.MajorOfferingImport) == 2


@pytest.mark.parametrize("index,value", [(0, 2036), (1, "不明地区"), (2, "abc"), (3, None), (5, " "), (6, None), (7, "=1+1")])
def test_invalid_row_rejects_entire_file(app, session, index, value):
    bad = ROW.copy()
    bad[index] = value
    result = upload(TestClient(app), excel([ROW, bad]))
    assert result.status_code == 422 and "第 3 行" in result.text
    assert count(session, models.MajorOfferingRecord) == 0
    assert count(session, models.MajorOfferingImport) == 0


def test_template_optional_columns_empty_and_invalid_files(app, session, monkeypatch):
    client = TestClient(app)
    response = client.get("/internal/v1/major-offerings/import-template")
    assert response.status_code == 200 and "spreadsheetml" in response.headers["content-type"]
    workbook = load_workbook(BytesIO(response.content))
    assert tuple(cell.value for cell in next(workbook.active.rows)) == HEADERS
    assert workbook.active["C2"].number_format == "@"
    workbook.close()
    assert upload(client, response.content).status_code == 422
    assert upload(client, b"not excel").status_code == 422
    assert upload(client, excel([ROW]), filename="file.xls").status_code == 422
    assert upload(client, excel([ROW], headers=["错误列"])).status_code == 422
    assert upload(client, excel([ROW]), filename="file.csv").status_code == 422
    assert upload(client, b"x" * (10 * 1024 * 1024 + 1)).status_code == 413
    monkeypatch.setattr("nexus_app.major_offering_import.MAX_ROWS", 1)
    assert upload(client, excel([ROW, ROW])).status_code == 422
    minimal = excel([[2026, "云南", 530701, "电子商务", "云南学院", 3]],
                    headers=["年份", "省份/地区", "专业代码", "专业名称", "学校名称", "修业年限"])
    assert upload(client, minimal).status_code == 200
    assert count(session, models.MajorOfferingRecord) == 1


def test_import_requires_console_role_and_key_and_is_not_public(app, app_no_auth_override, stub_user):
    assert upload(TestClient(app_no_auth_override), excel([ROW])).status_code == 401
    client = TestClient(app)
    assert client.post(PATH, files={"file": ("file.xlsx", excel([ROW]))}).status_code == 428
    stub_user.role = UserRole.BUSINESS_EXPERT
    assert upload(client, excel([ROW])).status_code == 200
    for role in (UserRole.OPS, UserRole.API_CALLER):
        stub_user.role = role
        assert upload(client, excel([ROW])).status_code == 403
    assert client.post("/v1/major-offerings/import").status_code == 404


def test_zero_formatted_numeric_school_code_is_preserved():
    workbook = Workbook()
    workbook.active.append(HEADERS)
    workbook.active.append([*ROW[:4], 123, *ROW[5:]])
    workbook.active["E2"].number_format = "000000"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    assert parse_excel(output.getvalue())[0][1]["school_code"] == "000123"
