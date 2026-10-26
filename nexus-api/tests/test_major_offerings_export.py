from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest
from sqlalchemy import select

from nexus_app import models
from test_major_offerings import seed


@pytest.mark.parametrize("path", ["/v1/major-offerings/export", "/internal/v1/major-offerings/export"])
def test_excel_exports_all_filtered_rows_with_text_codes_and_audit(app, session, path):
    seed(session)
    row = session.get(models.MajorOfferingRecord, "1")
    row.major_code = "0530701"
    row.school_code = "0012345"
    row.remark = "=HYPERLINK(\"https://example.test\")\x01"
    session.commit()
    result = TestClient(app).get(path, params={"year": 2026, "province_name": "北京", "school_name": "学校", "page": 2, "pageSize": 1})
    assert result.status_code == 200, result.text[:100]
    assert result.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "major-offerings.xlsx" in result.headers["content-disposition"]
    assert "filename*=UTF-8''" in result.headers["content-disposition"]
    workbook = load_workbook(BytesIO(result.content))
    sheet = workbook["专业布点"]
    rows = list(sheet.values)
    assert rows[0] == ("年份", "省份", "专业代码", "专业名称", "学校标识码", "学校名称", "年限", "备注")
    assert len(rows) == 3
    assert all(row[0] == 2026 and row[1] == "北京市" for row in rows[1:])
    first = next(row for row in rows[1:] if row[5] == "学校甲")
    assert first[2] == "0530701" and first[4] == "0012345"
    assert first[7].startswith("'=") and "\x01" not in first[7]
    assert all(cell.data_type != "f" for row in sheet for cell in row)
    audit = session.scalar(select(models.AuditLog).where(models.AuditLog.target_id == "export"))
    assert audit.target_type == "major_offering_record"
    assert audit.summary["result_count"] == 2 and audit.trace_id and audit.actor_id


def test_empty_export_contains_headers_only(app, session):
    seed(session)
    response = TestClient(app).get("/internal/v1/major-offerings/export?major_name=不存在")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.max_row == 1 and sheet.max_column == 8


@pytest.mark.parametrize("path", ["/v1/major-offerings/export", "/internal/v1/major-offerings/export"])
def test_export_requires_authentication(app_no_auth_override, path):
    assert TestClient(app_no_auth_override).get(path).status_code == 401
